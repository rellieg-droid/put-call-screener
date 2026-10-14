# -*- coding: utf-8 -*-
"""בדיקות offline על נתונים סינתטיים. ללא רשת."""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from rules import Rules, evaluate_strike, live_mid, PASS, FAIL, UNKNOWN
from screener import (screen_chain, scan, occ_symbol, nearest_expiry, COLUMNS,
                      results_frame, tally_frame)
from data_yf import delta_fn, xem_fn

TODAY = date(2026, 10, 14)
EXP = "2026-11-13"   # 30 ימים


def mk_chain(side, price=100.0):
    """שרשרת סינתטית: IV 40%, ציטוט חי."""
    rows = []
    for k in range(60, 141, 5):
        k = float(k)
        otm = (price - k) if side == "PUT" else (k - price)
        mid = max(0.05, 3.0 * (0.85 ** (max(otm, 0) / 2.5)))
        rows.append({"strike": k, "bid": round(mid * 0.95, 2), "ask": round(mid * 1.05, 2),
                     "lastPrice": mid, "impliedVolatility": 0.40,
                     "openInterest": 120, "volume": 10})
    return rows


def run(side, rules=None, **kw):
    rules = rules or Rules(check_earnings=False)
    return screen_chain("TST", side, 100.0, EXP, mk_chain(side), rules,
                        delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY, **kw)


def test_live_mid():
    assert live_mid(1.0, 1.2) == 1.1
    assert live_mid(0, 1.2) is None
    assert live_mid(None, 1.2) is None
    assert live_mid(1.3, 1.2) is None


def test_occ_symbol_put_and_call():
    assert occ_symbol("LLY", "2026-10-30", 790, "PUT") == "LLY261030P00790000"
    assert occ_symbol("LLY", "2026-10-30", 790, "CALL") == "LLY261030C00790000"


def test_put_only_strikes_below_price():
    rows = run("PUT")
    assert rows and all(r["strike"] <= 100.0 for r in rows)
    assert all(r["side"] == "PUT" for r in rows)


def test_call_only_strikes_above_price():
    rows = run("CALL")
    assert rows and all(r["strike"] >= 100.0 for r in rows)
    assert all(r["side"] == "CALL" for r in rows)


def test_delta_sign_and_symmetry():
    p = [r for r in run("PUT") if r["strike"] == 85.0][0]
    c = [r for r in run("CALL") if r["strike"] == 115.0][0]
    assert p["delta_signed"] < 0 < c["delta_signed"]
    assert 0 < p["delta_abs"] < 0.5 and 0 < c["delta_abs"] < 0.5


def test_xem_positive_for_otm_both_sides():
    assert all(r["xem"] is not None and r["xem"] >= 0 for r in run("PUT"))
    assert all(r["xem"] is not None and r["xem"] >= 0 for r in run("CALL"))


def test_delta_rule_filters():
    rules = Rules(delta_min=0.05, delta_max=0.10, check_earnings=False, check_exposure=False,
                  min_open_interest=0, max_spread_pct=0)
    for side in ("PUT", "CALL"):
        for r in run(side, rules):
            if not r["failed"]:
                assert 0.05 <= r["delta_abs"] <= 0.10


def test_yield_and_breakeven_put_vs_call():
    p = [r for r in run("PUT") if r["strike"] == 90.0][0]
    c = [r for r in run("CALL") if r["strike"] == 110.0][0]
    assert abs(p["breakeven"] - (90.0 - p["mid"])) < 1e-9
    assert abs(c["breakeven"] - (100.0 - c["mid"])) < 1e-9
    assert abs(p["annual_yield_pct"] - p["mid"] / 90.0 * (365 / 30) * 100) < 1e-6
    assert abs(c["annual_yield_pct"] - c["mid"] / 100.0 * (365 / 30) * 100) < 1e-6


def test_commitment_put_vs_call():
    p = [r for r in run("PUT", contracts=2) if r["strike"] == 90.0][0]
    c = [r for r in run("CALL", contracts=2) if r["strike"] == 110.0][0]
    assert p["commitment"] == 90.0 * 100 * 2
    assert c["commitment"] == 100.0 * 100 * 2


def test_exposure_rule_put_only():
    rules = Rules(delta_min=0, delta_max=0, check_earnings=False, check_exposure=True,
                  min_open_interest=0, max_spread_pct=0)
    put = run("PUT", rules, settled_cash=10000, alloc_limit=70.0)
    call = run("CALL", rules, settled_cash=10000, alloc_limit=70.0)
    assert any("חשיפה" in r["results"] for r in put)
    assert all("חשיפה" not in r["results"] for r in call)
    over = [r for r in put if r["commitment"] / 10000 * 100 > 70.0]
    assert over and all(r["failed"] for r in over)


def test_no_live_quote_fails_and_never_passes():
    ch = mk_chain("PUT")
    for c in ch:
        c["bid"] = 0.0
    rows = screen_chain("TST", "PUT", 100.0, EXP, ch, Rules(check_earnings=False),
                        delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY)
    assert rows and all(r["failed"] for r in rows)


def test_unknown_never_counts_as_pass():
    res = evaluate_strike({"side": "PUT", "mid": 1.0, "delta_abs": None, "open_interest": None,
                           "spread_pct": 1.0, "earnings_before_expiry": None},
                          Rules(check_exposure=False))
    assert res["Delta"][0] == UNKNOWN
    assert res["Open Interest"][0] == UNKNOWN
    assert res["דוחות"][0] == UNKNOWN


def test_earnings_rule():
    rules = Rules(delta_min=0, delta_max=0, min_open_interest=0, max_spread_pct=0,
                  check_exposure=False, check_earnings=True)
    inside = run("PUT", rules, earnings_date=date(2026, 11, 1), earnings_known=True)
    outside = run("PUT", rules, earnings_date=date(2026, 12, 20), earnings_known=True)
    assert all(r["results"]["דוחות"][0] == FAIL for r in inside)
    assert all(r["results"]["דוחות"][0] == PASS for r in outside)


def test_scan_both_sides_columns_and_sorting():
    rules = Rules(delta_min=0.05, delta_max=0.25, check_earnings=False, check_exposure=False,
                  min_open_interest=0, max_spread_pct=0)
    out = scan(["BBB", "AAA"], date(2026, 11, 13), rules, ["PUT", "CALL"],
               fetch_spot=lambda t: 100.0,
               fetch_expirations=lambda t: ["2026-10-30", EXP, "2026-12-18"],
               fetch_chain=lambda t, e, s: mk_chain(s),
               earnings_fn=lambda t: (False, None),
               delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY)
    rows = out["rows"]
    assert {r["Type"] for r in rows} == {"PUT", "CALL"}
    assert [r["Ticker"] for r in rows] == sorted(r["Ticker"] for r in rows)
    assert all(r["Expiration"] == EXP for r in rows)
    df = results_frame(rows)
    assert list(df.columns) == COLUMNS
    puts = df[df.Type == "PUT"]
    calls = df[df.Type == "CALL"]
    assert (puts.Delta < 0).all() and (calls.Delta > 0).all()
    assert len(out["summary"]) == 4


def test_scan_survives_ticker_failure():
    def spot(t):
        if t == "BAD":
            raise RuntimeError("boom")
        return 100.0
    rules = Rules(check_earnings=False, check_exposure=False)
    out = scan(["BAD", "OK"], date(2026, 11, 13), rules, ["PUT"],
               fetch_spot=spot, fetch_expirations=lambda t: [EXP],
               fetch_chain=lambda t, e, s: mk_chain(s),
               earnings_fn=lambda t: (False, None),
               delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY)
    notes = {s["ticker"]: s["note"] for s in out["summary"]}
    assert "RuntimeError" in notes["BAD"]
    ok = [x for x in out["summary"] if x["ticker"] == "OK"][0]
    assert ok["checked"] > 0   # OK נסרק למרות ש-BAD נכשל


def test_tally_frame():
    rules = Rules(check_earnings=False, check_exposure=False)
    out = scan(["AAA"], date(2026, 11, 13), rules, ["PUT"],
               fetch_spot=lambda t: 100.0, fetch_expirations=lambda t: [EXP],
               fetch_chain=lambda t, e, s: mk_chain(s),
               earnings_fn=lambda t: (False, None),
               delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY)
    df = tally_frame(out["tally"])
    assert "כלל" in df.columns and len(df) >= 1


def test_nearest_expiry():
    assert nearest_expiry(["2026-10-30", EXP, "2026-12-18"], date(2026, 11, 10)) == EXP
    assert nearest_expiry([], date(2026, 11, 10)) is None


def test_defaults_match_user_rules():
    r = Rules()
    assert (r.delta_min, r.delta_max, r.min_premium) == (0.05, 0.10, 1.0)


def test_min_premium_rule():
    rules = Rules(delta_min=0, delta_max=0, check_earnings=False, check_exposure=False,
                  min_open_interest=0, max_spread_pct=0, min_premium=1.0)
    for side in ("PUT", "CALL"):
        rows = run(side, rules)
        assert rows
        for r in rows:
            if r["mid"] is not None and r["mid"] < 1.0:
                assert r["results"]["פרמיה"][0] == FAIL and r["failed"]
            if not r["failed"]:
                assert r["mid"] >= 1.0
    off = Rules(delta_min=0, delta_max=0, check_earnings=False, check_exposure=False,
                min_open_interest=0, max_spread_pct=0, min_premium=0)
    assert all("פרמיה" not in r["results"] for r in run("PUT", off))


def test_sheet_columns_match_user_sheet_exactly():
    from screener import SHEET_COLUMNS, EXTRA_COLUMNS
    assert SHEET_COLUMNS == [
        "Ticker", "Type", "ContractSymbol", "Strike", "UnderlyingPrice", "Bid", "Ask", "LastPrice",
        "PriceUsed", "Delta", "IV_Percent", "DTE", "Expiration", "AnnualizedReturn_Percent",
        "RiskFreeRate_Percent", "OpenInterest", "Volume", "IsOffHours"]
    assert COLUMNS == SHEET_COLUMNS + EXTRA_COLUMNS


def _scan_df():
    rules = Rules(delta_min=0.03, delta_max=0.30, min_premium=0, check_earnings=False,
                  check_exposure=False, min_open_interest=0, max_spread_pct=0)
    out = scan(["AAA", "BBB"], date(2026, 11, 13), rules, ["PUT", "CALL"],
               fetch_spot=lambda t: 100.0, fetch_expirations=lambda t: [EXP],
               fetch_chain=lambda t, e, s: mk_chain(s), earnings_fn=lambda t: (False, None),
               delta_fn=delta_fn, xem_fn=xem_fn, today=TODAY, off_hours=True)
    return results_frame(out["rows"])


def test_excel_export_is_a_real_table_like_her_sheet():
    from io import BytesIO
    import datetime as dtm
    from openpyxl import load_workbook
    from screener import export_excel, SHEET_COLUMNS
    df = _scan_df()[SHEET_COLUMNS]
    wb = load_workbook(BytesIO(export_excel(df)))
    ws = wb.active
    assert [c.value for c in ws[1]] == SHEET_COLUMNS
    assert "Table1" in ws.tables
    assert ws.tables["Table1"].ref == f"A1:R{len(df) + 1}"
    assert ws.max_row == len(df) + 1
    ci = SHEET_COLUMNS.index("Expiration") + 1
    cell = ws.cell(row=2, column=ci)
    assert isinstance(cell.value, (dtm.date, dtm.datetime)) and cell.number_format == "m/d/yyyy"
    oi = SHEET_COLUMNS.index("IsOffHours") + 1
    assert ws.cell(row=2, column=oi).value is True
    ti = SHEET_COLUMNS.index("Type") + 1
    assert {ws.cell(row=r, column=ti).value for r in range(2, ws.max_row + 1)} == {"PUT", "CALL"}


def test_excel_export_empty_and_extended():
    from io import BytesIO
    from openpyxl import load_workbook
    from screener import export_excel
    empty = results_frame([])
    ws = load_workbook(BytesIO(export_excel(empty))).active
    assert ws.max_row == 1 and ws["A1"].value == "Ticker"
    ext = _scan_df()
    ws2 = load_workbook(BytesIO(export_excel(ext))).active
    assert [c.value for c in ws2[1]] == COLUMNS
