# -*- coding: utf-8 -*-
"""
screener.py

סריקת שרשראות אופציה ל-PUT ול-CALL מול הכללים של המשתמשת.
מחזיר שורה לכל חוזה שלא נפסל באף כלל, בפריסת העמודות של הגיליון שלה.

הנחות:
- PUT = מכירת PUT מכוסה במזומן (התחייבות = סטרייק x 100).
- CALL = מכירת CALL מכוסה (יש לה 100 מניות לכל חוזה), לא CALL חשוף.
  התחייבות = שווי 100 המניות במחיר הנוכחי. כלל החשיפה חל על PUT בלבד.

אין ציון ואין דירוג: הרשימה ממוינת לפי טיקר, סוג וסטרייק בלבד.
הנתונים והנוסחאות מוזרקים מבחוץ, כך שאפשר לבדוק offline.
"""
from __future__ import annotations

from datetime import date

import pandas as pd

from rules import Rules, evaluate_strike, num, live_mid, FAIL, PASS, RULE_EXPLAIN

MULT = 100

# העמודות של הגיליון של המשתמשת, בסדר הכותרות שלו (A עד R). זה מה שמוצג ומיוצא כברירת מחדל.
SHEET_COLUMNS = [
    "Ticker", "Type", "ContractSymbol", "Strike", "UnderlyingPrice", "Bid", "Ask",
    "LastPrice", "PriceUsed", "Delta", "IV_Percent", "DTE", "Expiration",
    "AnnualizedReturn_Percent", "RiskFreeRate_Percent", "OpenInterest", "Volume",
    "IsOffHours",
]
# עמודות נוספות, אופציונליות, אחרי עמודות הגיליון
EXTRA_COLUMNS = [
    "Distance_Percent", "xEM", "BreakEven", "Commitment_USD", "Spread_Percent",
    "CumExposure_Percent", "NotChecked",
]
COLUMNS = SHEET_COLUMNS + EXTRA_COLUMNS


def occ_symbol(ticker: str, expiry_iso: str, strike: float, side: str = "PUT") -> str:
    """OCC-style symbol, e.g. LLY261030P00790000."""
    y, m, d = expiry_iso.split("-")
    cp = "C" if side.upper() == "CALL" else "P"
    return f"{ticker}{y[2:]}{m}{d}{cp}{int(round(float(strike) * 1000)):08d}"


def nearest_expiry(expirations: list, target: date):
    if not expirations:
        return None
    return min(expirations, key=lambda e: abs((date.fromisoformat(e) - target).days))


def atm_iv(chain: list, price: float):
    """IV של הסטרייק הקרוב ביותר למחיר (עם IV תקין). None אם אין."""
    cands = []
    for r in chain:
        k, iv = num(r.get("strike")), num(r.get("impliedVolatility"))
        if k is not None and iv is not None and iv > 0.01:
            cands.append((abs(k - price), iv))
    return min(cands)[1] if cands else None


def screen_chain(ticker: str, side: str, price: float, expiry: str, chain: list, rules: Rules, *,
                 delta_fn, xem_fn, contracts: int = 1, settled_cash: float = 0.0,
                 alloc_limit: float | None = None, open_commitment: float = 0.0,
                 earnings_date: date | None = None, earnings_known: bool = False,
                 r: float = 0.045, today: date | None = None, span: float = 0.5) -> list[dict]:
    """
    כל חוזה OTM בטווח (PUT: [(1-span)*price, price]; CALL: [price, (1+span)*price])
    -> שורה עם נתונים גולמיים + תוצאת כל כלל.
    delta_fn(side, S, K, T_years, iv, r) -> delta | None
    xem_fn(side, S, K, sigma, days) -> distance in expected moves | None
    """
    side = side.upper()
    today = today or date.today()
    exp_d = date.fromisoformat(expiry)
    dte = max(1, (exp_d - today).days)
    T = dte / 365.0
    sigma = atm_iv(chain, price)
    if earnings_known:
        e_flag = bool(earnings_date is not None and today <= earnings_date <= exp_d)
    else:
        e_flag = None
    lo, hi = ((1 - span) * price, price) if side == "PUT" else (price, (1 + span) * price)

    rows = []
    for c in chain:
        k = num(c.get("strike"))
        if k is None or not (lo <= k <= hi):
            continue
        iv = num(c.get("impliedVolatility"))
        delta = None
        if iv and iv > 0:
            try:
                delta = delta_fn(side, price, k, T, iv, r)
            except Exception:
                delta = None
        xem = None
        if sigma:
            try:
                xem = xem_fn(side, price, k, sigma, dte)
            except Exception:
                xem = None
        mid = live_mid(c.get("bid"), c.get("ask"))
        if side == "PUT":
            commitment = k * MULT * contracts
            basis = k
            breakeven = None if mid is None else k - mid
            dist = (price - k) / price * 100.0
        else:
            commitment = price * MULT * contracts
            basis = price
            breakeven = None if mid is None else price - mid
            dist = (k - price) / price * 100.0
        row = {
            "ticker": ticker, "side": side, "expiry": expiry, "dte": dte, "price": price, "strike": k,
            "dist_pct": dist,
            "delta_abs": None if delta is None else abs(float(delta)),
            "delta_signed": None if delta is None else float(delta),
            "xem": None if xem is None else float(xem),
            "mid": mid,
            "spread_pct": None if mid is None else (c["ask"] - c["bid"]) / mid * 100.0,
            "breakeven": breakeven,
            "annual_yield_pct": None if mid is None else mid / basis * (365.0 / dte) * 100.0,
            "open_interest": num(c.get("openInterest")),
            "volume": num(c.get("volume")),
            "commitment": commitment,
            "earnings_before_expiry": e_flag,
            "earnings_date": earnings_date,
            "alloc_limit": alloc_limit if side == "PUT" else None,
            "cum_pct_settled": ((open_commitment + commitment) / settled_cash * 100.0
                                if side == "PUT" and settled_cash and settled_cash > 0 else None),
        }
        res = evaluate_strike(row, rules)
        row["results"] = res
        row["failed"] = [v[1] for v in res.values() if v[0] == FAIL]
        row["unknown"] = [n for n, v in res.items() if v[0] == "unknown"]
        rows.append(row)
    rows.sort(key=lambda x: -x["strike"])
    return rows


def _r(v, nd):
    return None if v is None else round(float(v), nd)


def _int(v):
    return None if v is None else int(v)


def contract_row(srow: dict, chain_by_strike: dict, r: float, off_hours) -> dict:
    c = chain_by_strike.get(round(srow["strike"], 4), {})
    iv = num(c.get("impliedVolatility"))
    return {
        "Ticker": srow["ticker"],
        "Type": srow["side"],
        "ContractSymbol": occ_symbol(srow["ticker"], srow["expiry"], srow["strike"], srow["side"]),
        "Strike": srow["strike"],
        "UnderlyingPrice": _r(srow["price"], 2),
        "Bid": num(c.get("bid")),
        "Ask": num(c.get("ask")),
        "LastPrice": num(c.get("lastPrice")),
        "PriceUsed": _r(srow["mid"], 2),
        "Delta": None if srow["delta_signed"] is None else round(srow["delta_signed"], 4),
        "IV_Percent": None if not iv else round(iv * 100.0, 2),
        "DTE": srow["dte"],
        "Expiration": srow["expiry"],
        "AnnualizedReturn_Percent": _r(srow["annual_yield_pct"], 1),
        "RiskFreeRate_Percent": round(r * 100.0, 2),
        "OpenInterest": _int(srow["open_interest"]),
        "Volume": _int(srow["volume"]),
        "IsOffHours": off_hours,
        "Distance_Percent": _r(srow["dist_pct"], 1),
        "xEM": _r(srow["xem"], 2),
        "BreakEven": _r(srow["breakeven"], 2),
        "Commitment_USD": round(srow["commitment"], 0),
        "Spread_Percent": _r(srow["spread_pct"], 0),
        "CumExposure_Percent": _r(srow["cum_pct_settled"], 0),
        "NotChecked": ", ".join(srow["unknown"]) if srow["unknown"] else "",
    }


def scan(tickers: list[str], target_expiry: date, rules: Rules, sides: list[str], *,
         fetch_spot, fetch_expirations, fetch_chain, earnings_fn, delta_fn, xem_fn,
         contracts: int = 1, settled_cash: float = 0.0, alloc_limit: float | None = None,
         open_commitment: float = 0.0, r: float = 0.045, off_hours=None,
         progress=None, today: date | None = None) -> dict:
    """
    fetch_chain(ticker, expiry, side) -> list[dict]
    earnings_fn(ticker) -> (known, date|None)
    מחזיר {"rows": [...], "summary": [...], "tally": {...}}.
    כשל בטיקר אחד לא עוצר את הסריקה, הוא מדווח בהערה של אותו טיקר.
    """
    today = today or date.today()
    rows_out: list[dict] = []
    summary: list[dict] = []
    tally = {"strikes": 0, "fail": {}, "only": {}}
    seen: set[str] = set()
    todo = []
    for t in tickers:
        t = (t or "").strip().upper()
        if t and t not in seen:
            seen.add(t)
            todo.append(t)

    for i, t in enumerate(todo):
        if progress:
            try:
                progress(i, len(todo), t)
            except Exception:
                pass
        try:
            price = fetch_spot(t)
        except Exception as e:
            price = None
            err = type(e).__name__
        else:
            err = ""
        if price is None:
            for sd in sides:
                summary.append({"ticker": t, "side": sd, "expiry": None, "checked": 0, "passed": 0,
                                "note": "לא נמצא מחיר" + (f" ({err})" if err else "")})
            continue
        try:
            exp = nearest_expiry(fetch_expirations(t), target_expiry)
        except Exception as e:
            exp, err = None, type(e).__name__
        if exp is None:
            for sd in sides:
                summary.append({"ticker": t, "side": sd, "expiry": None, "checked": 0, "passed": 0,
                                "note": "אין תאריכי פקיעה" + (f" ({err})" if err else "")})
            continue
        known, ed = (False, None)
        if rules.check_earnings:
            try:
                known, ed = earnings_fn(t)
            except Exception:
                known, ed = (False, None)
        for sd in sides:
            entry = {"ticker": t, "side": sd, "expiry": exp, "checked": 0, "passed": 0, "note": ""}
            summary.append(entry)
            try:
                chain = fetch_chain(t, exp, sd)
                if not chain:
                    entry["note"] = f"אין שרשרת {sd}"
                    continue
                srows = screen_chain(
                    t, sd, float(price), exp, chain, rules, delta_fn=delta_fn, xem_fn=xem_fn,
                    contracts=int(contracts), settled_cash=settled_cash, alloc_limit=alloc_limit,
                    open_commitment=open_commitment, earnings_date=ed, earnings_known=known,
                    r=r, today=today)
                entry["checked"] = len(srows)
                by_strike = {round(float(c["strike"]), 4): c for c in chain if num(c.get("strike")) is not None}
                for sr in srows:
                    tally["strikes"] += 1
                    fr = [n for n, (st_, _) in sr["results"].items() if st_ == FAIL]
                    for n in fr:
                        tally["fail"][n] = tally["fail"].get(n, 0) + 1
                    if len(fr) == 1:
                        tally["only"][fr[0]] = tally["only"].get(fr[0], 0) + 1
                    if sr["failed"]:
                        continue
                    rows_out.append(contract_row(sr, by_strike, r, off_hours))
                    entry["passed"] += 1
                if not srows:
                    entry["note"] = "אין חוזים בטווח"
                elif rules.check_earnings and not known:
                    entry["note"] = "תאריך דוח לא ידוע"
            except Exception as e:
                entry["note"] = f"שגיאה: {type(e).__name__}"

    rows_out.sort(key=lambda x: (x["Ticker"], x["Type"], -x["Strike"]))
    return {"rows": rows_out, "summary": summary, "tally": tally}


def tally_frame(tally: dict) -> pd.DataFrame:
    """לכל כלל: כמה חוזים פסל וכמה נפסלו רק בגללו (רמז איזה כלל לרופף)."""
    out = [{"כלל": n, "חוזים שנפסלו": c, "נפסלו רק בגללו": tally["only"].get(n, 0),
            "מה זה אומר": RULE_EXPLAIN.get(n, "")}
           for n, c in sorted(tally["fail"].items(), key=lambda x: -x[1])]
    return pd.DataFrame(out)


def summary_frame(summary: list[dict]) -> pd.DataFrame:
    return pd.DataFrame([{
        "טיקר": s["ticker"], "סוג": s["side"], "פקיעה": s["expiry"] or "-",
        "נבדקו": s["checked"], "עברו": s["passed"], "הערה": s["note"],
    } for s in summary])


def results_frame(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows, columns=COLUMNS)
    return df


def export_excel(df: pd.DataFrame, sheet_name: str = "Sheet1") -> bytes:
    """
    קובץ xlsx עם טבלת Excel אמיתית (Table1, עם כפתורי סינון בכותרות ופסים), כמו בגיליון שלה.
    Expiration נכתב כתאריך אמיתי, IsOffHours כ-TRUE/FALSE, ערכים חסרים כתא ריק.
    """
    from datetime import date as _date
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    cols = list(df.columns)
    ws.append(cols)

    def _cell(col, v):
        if v is None or (isinstance(v, float) and v != v):
            return None
        if col == "Expiration":
            try:
                return _date.fromisoformat(str(v))
            except ValueError:
                return v
        if hasattr(v, "item"):          # numpy scalar -> python
            v = v.item()
        return v

    for rec in df.to_dict("records"):
        ws.append([_cell(c, rec[c]) for c in cols])

    n_rows = len(df) + 1
    if "Expiration" in cols:
        ci = cols.index("Expiration") + 1
        for r in range(2, n_rows + 1):
            ws.cell(row=r, column=ci).number_format = "m/d/yyyy"

    for i, c in enumerate(cols, start=1):
        width = max(len(str(c)) + 4, 11)
        if c == "ContractSymbol":
            width = 24
        ws.column_dimensions[get_column_letter(i)].width = width

    ref = f"A1:{get_column_letter(len(cols))}{max(n_rows, 2)}"
    tbl = Table(displayName="Table1", ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                        showColumnStripes=False)
    ws.add_table(tbl)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
